"""Run entity_eval on the dataset with parallel workers and failed-sample retry."""

import argparse
import json
import os
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import dataclass

import pandas as pd

from entity_eval.agent import EntityEvalAgent, SYSTEM_PROMPT_RELAXED
from reviewer.agent import ReviewerAgent


@dataclass
class ExtractionFailureRecord:
    doc_id: str
    success: bool = False
    duration_seconds: float = 0
    prompt_tokens: int = 0
    completion_tokens: int = 0
    entities_count: int = 0
    error: str = ""


ENTITY_COLUMNS = [
    "doc_id",
    "entity",
    "mapped_from",
    "entity_sentiment",
    "impact_level",
    "sentiment_reason",
    "risk_type",
]
OUTPUT_FORMATS = {"csv", "jsonl", "xlsx"}

IMPACT_RANK = {"": 0, "小": 1, "中": 2, "大": 3}
FILTER_REASON_PATTERNS = [
    "仅提及",
    "仅作为",
    "仅被提及",
    "无具体事件或数据",
    "无具体事件或业务影响",
    "无具体协议或业务影响",
    "没有具体事件",
    "未提及具体",
    "未涉及其自身",
    "只是背景",
    "作为对比",
    "历史背景",
    "成分股",
    "成交额",
]


def _dedupe_keep_order(values: list[str]) -> list[str]:
    seen = set()
    result = []
    for value in values:
        value = value.strip()
        if value and value not in seen:
            seen.add(value)
            result.append(value)
    return result


def _split_mapped_from(value: str) -> list[str]:
    if not value:
        return []
    return [part.strip() for part in value.replace("，", ",").split(",") if part.strip()]


def _merge_entity_rows(rows: list[dict]) -> list[dict]:
    grouped: dict[tuple[str, str, str], dict] = {}

    for row in rows:
        key = (row["doc_id"], row["entity"], row["entity_sentiment"])
        if key not in grouped:
            grouped[key] = {
                **row,
                "_mapped_from_parts": [],
                "_sentiment_reasons": [],
                "_risk_type_set": set(),
            }

        merged = grouped[key]
        merged["_mapped_from_parts"].extend(_split_mapped_from(row.get("mapped_from", "")))
        if row.get("sentiment_reason"):
            merged["_sentiment_reasons"].append(row["sentiment_reason"])

        current_level = merged.get("impact_level", "") or ""
        new_level = row.get("impact_level", "") or ""
        if IMPACT_RANK.get(new_level, 0) > IMPACT_RANK.get(current_level, 0):
            merged["impact_level"] = new_level

        raw_risk = row.get("risk_type", "") or ""
        if raw_risk:
            for part in raw_risk.replace("，", ",").split(","):
                part = part.strip()
                if part:
                    merged["_risk_type_set"].add(part)

    merged_rows = []
    for row in grouped.values():
        mapped_from = "，".join(_dedupe_keep_order(row.pop("_mapped_from_parts")))
        reasons = "；".join(_dedupe_keep_order(row.pop("_sentiment_reasons")))
        risk_set = row.pop("_risk_type_set")
        row["mapped_from"] = mapped_from
        row["sentiment_reason"] = reasons
        row["risk_type"] = "，".join(sorted(risk_set)) if risk_set else ""
        merged_rows.append(row)

    return merged_rows


def _filter_entity_rows(rows: list[dict]) -> tuple[list[dict], list[dict]]:
    kept = []
    filtered = []

    for row in rows:
        reason = row.get("sentiment_reason", "") or ""
        matched_pattern = next(
            (pattern for pattern in FILTER_REASON_PATTERNS if pattern in reason),
            "",
        )
        if matched_pattern:
            filtered.append({**row, "filter_reason": matched_pattern})
        else:
            kept.append(row)

    return kept, filtered


def _dedup_reviewer_entities(existing: list[dict], reviewer: list[dict]) -> list[dict]:
    seen = {(r["doc_id"], r["entity"]) for r in existing}
    result = list(existing)
    for r in reviewer:
        key = (r["doc_id"], r["entity"])
        if key not in seen:
            seen.add(key)
            result.append(r)
    return result


def _build_raw_lines(doc_ids: list[str], entity_rows: list[dict]) -> list[dict]:
    rows_by_doc: dict[str, list[dict]] = {doc_id: [] for doc_id in doc_ids}
    for row in entity_rows:
        doc_id = row["doc_id"]
        entity_payload = {key: row.get(key, "") for key in ENTITY_COLUMNS if key != "doc_id"}
        rows_by_doc.setdefault(doc_id, []).append(entity_payload)
    return [{"doc_id": doc_id, "entities": rows_by_doc.get(doc_id, [])} for doc_id in doc_ids]


def _parse_output_formats(value: str, allowed: set[str]) -> set[str]:
    formats = {item.strip().lower() for item in value.split(",") if item.strip()}
    invalid = formats - allowed
    if invalid:
        raise ValueError(f"unsupported output format(s): {', '.join(sorted(invalid))}")
    if not formats:
        raise ValueError("at least one output format is required")
    return formats


def _write_xlsx(df: pd.DataFrame, path: str, sheet_name: str,
                source_data: dict | None = None, doc_ids: list[str] | None = None):
    try:
        import openpyxl  # noqa: F401
    except ImportError as e:
        raise RuntimeError("XLSX output requires openpyxl. Install it with: uv pip install openpyxl") from e

    df.to_excel(path, index=False, sheet_name=sheet_name, engine="openpyxl")

    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter
    from openpyxl.styles import Alignment

    workbook = load_workbook(path)
    worksheet = workbook[sheet_name]
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions

    for column_cells in worksheet.columns:
        header = str(column_cells[0].value or "")
        max_len = max(len(str(cell.value or "")) for cell in column_cells[:200])
        column_width = min(max(max_len + 2, len(header) + 2), 60)
        worksheet.column_dimensions[column_cells[0].column_letter].width = column_width

    if source_data is not None and doc_ids is not None:
        _add_review_sheet(workbook, df, source_data, doc_ids)

    workbook.save(path)


def _add_review_sheet(workbook, df: pd.DataFrame, source_data: dict, doc_ids: list[str]):
    from openpyxl.styles import Alignment
    from openpyxl.utils import get_column_letter

    review_ws = workbook.create_sheet("review")
    headers = [
        "doc_id", "headline", "content",
        "entity", "mapped_from", "entity_sentiment",
        "impact_level", "sentiment_reason", "risk_type",
    ]
    review_ws.append(headers)

    entities_by_doc: dict[str, list[dict]] = {}
    for _, row in df.iterrows():
        entities_by_doc.setdefault(row["doc_id"], []).append(row.to_dict())

    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    row_idx = 2

    for doc_id in doc_ids:
        headline, content = source_data.get(doc_id, ("", ""))
        entity_rows = entities_by_doc.get(doc_id, [])

        if not entity_rows:
            review_ws.append([doc_id, headline, content, "", "", "", "", "", ""])
            row_idx += 1
        else:
            n = len(entity_rows)
            start_row = row_idx
            for i, ent in enumerate(entity_rows):
                review_ws.append([
                    doc_id if i == 0 else None,
                    headline if i == 0 else None,
                    content if i == 0 else None,
                    ent.get("entity", ""),
                    ent.get("mapped_from", ""),
                    ent.get("entity_sentiment", ""),
                    ent.get("impact_level", ""),
                    ent.get("sentiment_reason", ""),
                    ent.get("risk_type", ""),
                ])
                row_idx += 1

            if n > 1:
                for col in range(1, 4):
                    col_letter = get_column_letter(col)
                    review_ws.merge_cells(
                        f"{col_letter}{start_row}:{col_letter}{start_row + n - 1}"
                    )

    for row in review_ws.iter_rows(min_row=2, max_col=9):
        for cell in row:
            cell.alignment = center

    col_widths = {"A": 14, "B": 30, "C": 40, "D": 18, "E": 18, "F": 12, "G": 8, "H": 40, "I": 18}
    for col_letter, width in col_widths.items():
        review_ws.column_dimensions[col_letter].width = width

    review_ws.auto_filter.ref = review_ws.dimensions
    review_ws.freeze_panes = "A2"


def parse_args():
    parser = argparse.ArgumentParser()
    parser.add_argument("--batch", "-n", type=int, default=10, help="number of articles to process")
    parser.add_argument(
        "--sample-mode",
        choices=["head", "random"],
        default="head",
        help="sample selection mode: head keeps current behavior; random samples rows",
    )
    parser.add_argument("--seed", type=int, default=42, help="random seed when --sample-mode=random")
    parser.add_argument("--workers", "-w", type=int, default=2, help="parallel worker count")
    parser.add_argument("--output", "-o", default="output", help="output directory")
    parser.add_argument(
        "--output-formats",
        default="csv,jsonl",
        help="comma-separated output formats: csv,jsonl,xlsx (summary.json is always written)",
    )
    parser.add_argument("--data", default="data/test_news_0603.csv", help="input CSV path")
    parser.add_argument("--model", default=None, help="model name (default: LLM_MODEL env or deepseek-v4-flash)")
    parser.add_argument("--base-url", default=None, help="API base URL (default: LLM_BASE_URL env or https://opencode.ai/zen/go/v1)")
    rerun_group = parser.add_mutually_exclusive_group()
    rerun_group.add_argument("--rerun-empty", action="store_true",
                             help="rerun 0-entity articles with relaxed prompt for recall diagnosis")
    rerun_group.add_argument("--retry-empty", action="store_true",
                             help="rerun 0-entity articles with formal prompt and merge results back")
    parser.add_argument("--review", nargs="?", const="all", default=None,
                        choices=["all", "zero"],
                        help="run reviewer stage: 'all' (default) reviews every article, "
                             "'zero' only reviews articles with 0 extracted entities")
    return parser.parse_args()


def main(args):
    os.makedirs(args.output, exist_ok=True)
    output_formats = _parse_output_formats(args.output_formats, OUTPUT_FORMATS)

    df = pd.read_csv(args.data)
    if args.sample_mode == "random":
        batch_size = min(args.batch, len(df))
        batch = df.sample(n=batch_size, random_state=args.seed).copy()
    else:
        batch = df.head(args.batch).copy()
    source = {row["doc_id"]: (row["headline"], str(row["content"])) for _, row in batch.iterrows()}

    agent = EntityEvalAgent(model=args.model, base_url=args.base_url)

    def run_pass(doc_ids: list[str], pass_label: str, pool_size: int,
                  system_prompt: str | None = None):
        done_count = 0
        total = len(doc_ids)
        pass_records: list = []
        pass_entities: list[dict] = []
        pass_raw: list[dict] = []
        failed_ids: list[str] = []

        with ThreadPoolExecutor(max_workers=pool_size) as pool:
            meta = {}
            for doc_id in doc_ids:
                headline, content = source[doc_id]
                future = pool.submit(agent.evaluate, doc_id, headline, content,
                                     system_prompt=system_prompt)
                meta[future] = doc_id

            for future in as_completed(meta):
                doc_id = meta[future]
                headline = source[doc_id][0]
                done_count += 1
                snippet = headline[:50] + ("..." if len(headline) > 50 else "")
                try:
                    entities, record = future.result()
                    pass_records.append(record)

                    article_entities = []
                    for ent in entities:
                        d = {
                            "entity": ent.entity,
                            "mapped_from": ent.mapped_from,
                            "entity_sentiment": ent.entity_sentiment,
                            "impact_level": ent.impact_level,
                            "sentiment_reason": ent.sentiment_reason,
                            "risk_type": "，".join(ent.risk_type),
                        }
                        pass_entities.append({"doc_id": doc_id, **d})
                        article_entities.append(d)

                    pass_raw.append({"doc_id": doc_id, "entities": article_entities})

                    print(f"[{pass_label}] [{done_count:>3}/{total}] [{record.duration_seconds:5.1f}s] "
                          f"{record.prompt_tokens:>4} + {record.completion_tokens:>3} tok "
                          f"{record.entities_count} ent | {snippet}")

                except Exception as e:
                    pass_records.append(ExtractionFailureRecord(doc_id=doc_id, error=str(e)))
                    failed_ids.append(doc_id)
                    print(f"[{pass_label}] [{done_count:>3}/{total}] [  err  ] {e} | {snippet}")

        return pass_records, pass_entities, pass_raw, failed_ids

    t_start = time.time()

    records, all_entities, raw_lines, failed = run_pass(
        list(source.keys()), "1st", args.workers
    )

    retry_count = 0
    if failed:
        print(f"\n--- Retrying {len(failed)} failed samples ---")
        retry_records, retry_entities, retry_raw, still_failed = run_pass(
            failed, "retry", 1
        )
        records.extend(retry_records)
        all_entities.extend(retry_entities)
        raw_lines.extend(retry_raw)
        retry_count = len(failed) - len(still_failed)

    wall_clock = time.time() - t_start
    total_entities_before_filter = len(all_entities)
    all_entities, filtered_entities = _filter_entity_rows(all_entities)
    if filtered_entities:
        print(f"\n--- Post-filtered {len(filtered_entities)} entity rows ---")
        for row in filtered_entities:
            reason = row.get("sentiment_reason", "")
            snippet = reason[:90] + ("..." if len(reason) > 90 else "")
            print(f"[filter] {row['doc_id']} | {row['entity']} | "
                  f"{row['filter_reason']} | {snippet}")

    total_entities_before_merge = len(all_entities)
    all_entities = _merge_entity_rows(all_entities)

    retry_empty_recovered = 0
    retry_empty_articles = 0
    zero_entity_ids = [
        r.doc_id for r in records
        if r.success and getattr(r, "entities_count", 0) == 0
    ]

    if args.retry_empty and zero_entity_ids:
        print(f"\n--- Retrying {len(zero_entity_ids)} 0-entity articles with formal prompt ---")
        retry_records, retry_entities, retry_raw, _ = run_pass(
            zero_entity_ids, "retry-empty", min(2, args.workers),
        )
        retry_entities, _ = _filter_entity_rows(retry_entities)
        retry_entities = _merge_entity_rows(retry_entities)
        recovered_docs = set()
        for ent in retry_entities:
            all_entities.append(ent)
            recovered_docs.add(ent["doc_id"])
        retry_empty_recovered = len(retry_entities)
        retry_empty_articles = len(recovered_docs)
        records.extend(retry_records)
        wall_clock = time.time() - t_start
        print(f"  Recovered {retry_empty_recovered} entities from {retry_empty_articles} articles")

    reviewer_recovered = 0
    reviewer_recovered_articles = set()
    n_before_review = len(all_entities)
    if args.review:
        zero_entity_ids = [
            r.doc_id for r in records
            if r.success and getattr(r, "entities_count", 0) == 0
        ]
        if args.review == "zero":
            review_doc_ids = zero_entity_ids
            scope_desc = f"{len(review_doc_ids)} 0-entity articles"
        else:
            review_doc_ids = list(source.keys())
            scope_desc = f"{len(review_doc_ids)} articles"
        print(f"\n--- Review stage ({args.review}): scanning {scope_desc} ---")
        reviewer_agent = ReviewerAgent(model=args.model, base_url=args.base_url)
        existing_by_doc = {}
        for ent in all_entities:
            existing_by_doc.setdefault(ent["doc_id"], []).append({
                "entity": ent["entity"],
                "mapped_from": ent.get("mapped_from", ""),
                "entity_sentiment": ent.get("entity_sentiment", ""),
                "sentiment_reason": ent.get("sentiment_reason", ""),
            })

        reviewer_entities: list[dict] = []
        with ThreadPoolExecutor(max_workers=min(4, args.workers)) as pool:
            meta = {}
            for doc_id in review_doc_ids:
                existing = existing_by_doc.get(doc_id, [])
                headline, content = source[doc_id]
                future = pool.submit(reviewer_agent.review, doc_id, headline, content, existing)
                meta[future] = doc_id

            done_count = 0
            total = len(meta)
            reviewer_records: list = []
            for future in as_completed(meta):
                doc_id = meta[future]
                done_count += 1
                snippet = source[doc_id][0][:50] + ("..." if len(source[doc_id][0]) > 50 else "")
                try:
                    entities, record = future.result()
                    reviewer_records.append(record)
                    for ent in entities:
                        reviewer_entities.append({
                            "doc_id": doc_id,
                            "entity": ent.entity,
                            "mapped_from": ent.mapped_from,
                            "entity_sentiment": ent.entity_sentiment,
                            "impact_level": ent.impact_level,
                            "sentiment_reason": ent.sentiment_reason,
                            "risk_type": "，".join(ent.risk_type),
                        })
                    if record.entities_count > 0:
                        names = [e.entity for e in entities]
                        print(f"  [review] [{done_count:>3}/{total}] found {record.entities_count}: {names}")
                    else:
                        print(f"  [review] [{done_count:>3}/{total}] none | {snippet}")
                except Exception as e:
                    print(f"  [review] [{done_count:>3}/{total}] err: {e} | {snippet}")

        reviewer_entities, review_filtered = _filter_entity_rows(reviewer_entities)
        all_entities = _dedup_reviewer_entities(all_entities, reviewer_entities)
        reviewer_recovered = len(all_entities) - n_before_review
        reviewer_recovered_articles = {r["doc_id"] for r in all_entities[n_before_review:]}
        print(f"  Reviewer recovered {reviewer_recovered} entities "
              f"({len(reviewer_recovered_articles)} articles), "
              f"{len(review_filtered)} filtered by post-processing")
        records.extend(reviewer_records)
        wall_clock = time.time() - t_start

    raw_lines = _build_raw_lines(list(source.keys()), all_entities)
    entities_df = pd.DataFrame(all_entities, columns=ENTITY_COLUMNS)

    if "jsonl" in output_formats:
        with open(f"{args.output}/entities.jsonl", "w", encoding="utf-8") as f:
            for line in raw_lines:
                f.write(json.dumps(line, ensure_ascii=False) + "\n")

    if "csv" in output_formats:
        entities_df.to_csv(f"{args.output}/entities.csv", index=False, encoding="utf-8-sig")

    if "xlsx" in output_formats:
        _write_xlsx(entities_df, f"{args.output}/entities.xlsx", "entities",
                    source_data=source, doc_ids=list(source.keys()))

    if args.rerun_empty and zero_entity_ids:
        print(f"\n--- Rerunning {len(zero_entity_ids)} 0-entity articles with relaxed prompt ---")
        _, rerun_entities, rerun_raw, _ = run_pass(
            zero_entity_ids, "rerun", min(2, args.workers),
            system_prompt=SYSTEM_PROMPT_RELAXED,
        )
        rerun_entities = _merge_entity_rows(rerun_entities)
        rerun_df = pd.DataFrame(rerun_entities, columns=ENTITY_COLUMNS)
        rerun_df.to_csv(f"{args.output}/rerun_entities.csv", index=False, encoding="utf-8-sig")
        print(f"  Rerun extracted {len(rerun_entities)} entities across "
              f"{rerun_df['doc_id'].nunique()} articles")
        print(f"  See {args.output}/rerun_entities.csv for details")

    successful = [r for r in records if r.success]
    summary = {
        "config": {
            "model": agent.model,
            "max_workers": args.workers,
            "batch": args.batch,
            "sample_mode": args.sample_mode,
            "seed": args.seed if args.sample_mode == "random" else None,
            "output_formats": sorted(output_formats),
        },
        "wall_clock_seconds": round(wall_clock, 1),
        "total_calls": len(records),
        "first_pass": len(source),
        "retried": len(failed) if failed else 0,
        "retry_succeeded": retry_count,
        "success": sum(1 for r in records if r.success),
        "failed": sum(1 for r in records if not r.success),
        "total_prompt_tokens": sum(r.prompt_tokens for r in successful),
        "total_completion_tokens": sum(r.completion_tokens for r in successful),
        "avg_duration_seconds": round(
            sum(r.duration_seconds for r in successful) / len(successful), 2
        ) if successful else 0,
        "total_entities_before_filter": total_entities_before_filter,
        "post_filtered_entities": len(filtered_entities),
        "post_filter_details": [
            {
                "doc_id": row["doc_id"],
                "entity": row["entity"],
                "entity_sentiment": row["entity_sentiment"],
                "filter_reason": row["filter_reason"],
                "sentiment_reason": row.get("sentiment_reason", ""),
            }
            for row in filtered_entities
        ],
        "retry_empty_articles": retry_empty_articles,
        "retry_empty_recovered": retry_empty_recovered,
        "reviewer_recovered": reviewer_recovered,
        "reviewer_recovered_articles": len(reviewer_recovered_articles),
        "total_entities_before_merge": total_entities_before_merge,
        "total_entities": len(all_entities),
        "per_call": [
            {
                "doc_id": r.doc_id,
                "success": r.success,
                "duration_seconds": r.duration_seconds,
                "prompt_tokens": r.prompt_tokens,
                "completion_tokens": r.completion_tokens,
                "entities_count": r.entities_count,
                "error": getattr(r, "error", ""),
            }
            for r in records
        ],
    }
    with open(f"{args.output}/summary.json", "w", encoding="utf-8") as f:
        json.dump(summary, f, ensure_ascii=False, indent=2)

    print(f"\n{'='*50}")
    print(f"Done. {summary['success']}/{summary['total_calls']} calls succeeded "
          f"(retry saved {summary['retry_succeeded']}).")
    print(f"Wall clock: {wall_clock:.1f}s | Avg: {summary['avg_duration_seconds']}s/call")
    print(f"Tokens: {summary['total_prompt_tokens']} in + {summary['total_completion_tokens']} out")
    print(f"Entities extracted: {summary['total_entities']}")
    if args.retry_empty and retry_empty_recovered:
        print(f"Retry-empty recovered {retry_empty_recovered} entities from {retry_empty_articles} articles")
    if args.review and reviewer_recovered:
        print(f"Reviewer ({args.review}) recovered {reviewer_recovered} entities from {len(reviewer_recovered_articles)} articles")
    print(f"Output formats: {', '.join(summary['config']['output_formats'])}")
    print(f"Output: {args.output}/")


if __name__ == "__main__":
    main(parse_args())
